"""从语料 Excel 自动生成 ground truth JSON, 含三业务线 + 精确/模糊分类."""
import json
import random
import sys
from pathlib import Path
import openpyxl

ROOT = Path(__file__).resolve().parent.parent
random.seed(42)

FILES = [
    ("data/raw/ccgp_qa.xlsx", "bidding"),
    ("data/raw/enterprise_qa.xlsx", "enterprise"),
    ("data/raw/regulation_qa.xlsx", "regulation"),
]


def read_qa(path: Path, business_line: str) -> list[dict]:
    wb = openpyxl.load_workbook(path, read_only=True)
    ws = wb.active
    headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    qa = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        d = dict(zip(headers, row))
        q = str(d.get("question", "") or "").strip()
        a = str(d.get("answer", "") or "").strip()
        if not q or not a:
            continue
        qa.append({
            "question": q,
            "answer": a,
            "business_line": business_line,
            "section_title": str(d.get("section_title", "") or ""),
        })
    wb.close()
    return qa


def make_fuzzy(q: str) -> str:
    """生成模糊/变体问题."""
    variants = []
    # 1. 去掉末尾标点
    if q.endswith(("？", "?", "。", ".")):
        variants.append(q.rstrip("？?。."))
    # 2. 去掉开头的"什么是"/"什么情况下"
    for prefix in ("什么是", "什么情况下", "什么是"):
        if q.startswith(prefix):
            variants.append(q[len(prefix):])
            break
    # 3. 截取前 60% 作为模糊查询
    cut = max(10, int(len(q) * 0.6))
    if cut < len(q):
        variants.append(q[:cut])
    # 4. 同义替换
    replacements = {
        "什么时候": "何时",
        "在哪里": "何地",
        "是多少": "多少",
        "包括哪些": "有哪些",
        "应当如何": "怎么",
    }
    modified = q
    for old, new in replacements.items():
        if old in modified:
            modified = modified.replace(old, new, 1)
            break
    if modified != q:
        variants.append(modified)
    return random.choice(variants) if variants else q


def main():
    all_gt = []
    for rel_path, bl in FILES:
        path = ROOT / rel_path
        if not path.exists():
            print(f"SKIP: {path} not found")
            continue
        qa = read_qa(path, bl)
        print(f"{path.name}: {len(qa)} 条 ({bl})")

        for i, item in enumerate(qa):
            # 精确匹配
            all_gt.append({
                "question": item["question"],
                "gt_question": item["question"],
                "business_line": bl,
                "question_type": "exact",
            })
            # 模糊匹配 (30% 采样)
            if random.random() < 0.3:
                fuzzy_q = make_fuzzy(item["question"])
                if fuzzy_q != item["question"]:
                    all_gt.append({
                        "question": fuzzy_q,
                        "gt_question": item["question"],
                        "business_line": bl,
                        "question_type": "fuzzy",
                    })

    random.shuffle(all_gt)
    out_path = ROOT / "eval" / "ground_truth.json"
    out_path.parent.mkdir(exist_ok=True)
    out_path.write_text(
        json.dumps(all_gt, ensure_ascii=False, indent=2),
        encoding="utf-8"
    )
    print(f"\n总计: {len(all_gt)} 条 ground truth")
    # 统计
    from collections import Counter
    bl_counts = Counter(g["business_line"] for g in all_gt)
    qt_counts = Counter(g["question_type"] for g in all_gt)
    print(f"  业务线: {dict(bl_counts)}")
    print(f"  问题类型: {dict(qt_counts)}")
    print(f"  输出: {out_path}")


if __name__ == "__main__":
    main()
