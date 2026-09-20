# -*- coding: utf-8 -*-
"""生成《招投标 RAG 智能问答系统》介绍 PPT (16:9).

内容: 系统架构 / 技术栈 / RAG 与 Agent 算法 / 前端页面截图.
用法: .venv\\Scripts\\python.exe scripts\\gen_ppt.py
输出: 招投标RAG智能问答系统介绍.pptx (项目根目录)
"""
from __future__ import annotations

import sys
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
SHOT = ROOT / "tests" / "acceptance" / "screenshots"
OUT = ROOT / "招投标RAG智能问答系统介绍.pptx"

SLIDE_W, SLIDE_H = Inches(13.333), Inches(7.5)

# 主题色
PRIMARY = RGBColor(0x1F, 0x4E, 0x79)   # 深蓝
TEAL = RGBColor(0x31, 0x85, 0x9C)      # 青
LIGHT = RGBColor(0xDC, 0xE6, 0xF1)     # 浅蓝底
GOLD = RGBColor(0xBF, 0x8F, 0x00)      # 金
GREEN = RGBColor(0x4F, 0x8A, 0x5B)     # 绿
GRAY = RGBColor(0x59, 0x59, 0x59)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = "微软雅黑"


# ---------- 基础工具 ----------

def _set_text(tf, lines, size=14, color=None, bold=False, align=PP_ALIGN.LEFT,
              space_after=4):
    """lines: list[str | (str, dict)]"""
    tf.word_wrap = True
    first = True
    for item in lines:
        text, opt = (item, {}) if isinstance(item, str) else item
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = opt.get("align", align)
        p.space_after = Pt(opt.get("space_after", space_after))
        r = p.add_run()
        r.text = text
        f = r.font
        f.name = FONT
        f.size = Pt(opt.get("size", size))
        f.bold = opt.get("bold", bold)
        f.color.rgb = opt.get("color", color or GRAY)


def add_box(slide, x, y, w, h, fill=LIGHT, line=None, radius=True):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    sp = slide.shapes.add_shape(shape_type, x, y, w, h)
    sp.fill.solid()
    sp.fill.fore_color.rgb = fill
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line
        sp.line.width = Pt(0.75)
    sp.shadow.inherit = False
    return sp


def add_text(slide, x, y, w, h, lines, **kw):
    tb = slide.shapes.add_textbox(x, y, w, h)
    _set_text(tb.text_frame, lines, **kw)
    return tb


def slide_base(prs, title=None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
    if title:
        add_box(slide, 0, 0, SLIDE_W, Inches(0.9), fill=PRIMARY, radius=False)
        add_text(slide, Inches(0.5), Inches(0.12), Inches(12.3), Inches(0.66),
                 [title], size=24, color=WHITE, bold=True)
        add_box(slide, 0, Inches(0.9), SLIDE_W, Pt(3), fill=GOLD, radius=False)
    return slide


def add_image_fit(slide, path, x, y, max_w, max_h, border=True):
    with Image.open(path) as im:
        iw, ih = im.size
    w = int(iw * 9525)  # px → EMU
    h = int(ih * 9525)
    scale = min(max_w / w, max_h / h)
    w, h = int(w * scale), int(h * scale)
    px = x + int((max_w - w) / 2)
    py = y + int((max_h - h) / 2)
    pic = slide.shapes.add_picture(str(path), px, py, w, h)
    if border:
        pic.line.color.rgb = RGBColor(0xBD, 0xBD, 0xBD)
        pic.line.width = Pt(0.75)
    return pic, px, py, w, h


def add_arrow(slide, x1, y1, x2, y2, color=GOLD, width=1.5):
    """带箭头直线连接线 (tailEnd 箭头需写 XML)."""
    conn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x1, y1, x2, y2)
    conn.line.color.rgb = color
    conn.line.width = Pt(width)
    ln = conn.line._get_or_add_ln()
    tail = ln.makeelement(qn("a:tailEnd"), {"type": "arrow", "w": "med", "len": "med"})
    ln.append(tail)
    return conn


# ---------- 各页 ----------

def s_cover(prs):
    slide = slide_base(prs)
    add_box(slide, 0, 0, SLIDE_W, SLIDE_H, fill=PRIMARY, radius=False)
    add_box(slide, 0, Inches(4.05), SLIDE_W, Pt(4), fill=GOLD, radius=False)
    add_text(slide, Inches(1), Inches(2.15), Inches(11.3), Inches(1.4),
             [("招投标 RAG 智能问答系统", {"size": 44, "bold": True, "color": WHITE})])
    add_text(slide, Inches(1), Inches(3.35), Inches(11.3), Inches(0.6),
             [("检索增强生成 · ReAct 智能体 · 多专家协作 · 有依据地回答",
               {"size": 18, "color": RGBColor(0xBD, 0xD7, 0xEE)})])
    add_text(slide, Inches(1), Inches(4.45), Inches(11.3), Inches(1.6), [
        ("系统架构  |  技术栈  |  核心算法  |  前端展示", {"size": 16, "color": WHITE}),
        ("技术亮点: 检索 HitRate@5=100% · nDCG@20=0.9846 · 验收 60/60 · 硬闸门防幻觉",
         {"size": 13, "color": RGBColor(0xBD, 0xD7, 0xEE)}),
    ])


def s_arch(prs):
    slide = slide_base(prs, "一、系统架构 — 五层分层设计")
    rows = [
        ("前端层", "Next.js 14 + React 18", "对话问答 · Agent 过程可视化 · 评测看板 · 文档管理 · 知识图谱 · 企业资料", TEAL),
        ("API 层", "FastAPI (:8001)", "REST + SSE 流式 · JWT/RBAC 鉴权 · 限流 · 审计日志 · /docs", PRIMARY),
        ("Agent 编排层", "ReAct + 多专家 + 意图路由 + 证据门", "意图路由(三业务线裁剪工具) → ReAct 循环(≤3步) / 多专家并行(法规·案例·价格) → 证据门(无引用不生成)", GOLD),
        ("RAG 检索层", "解析 → 切片 → Query 规划 → 多路召回 → 精排", "页感知解析(OCR) · 语义切片(标题加权) · 指代消解+受控变体 · Dense+BM25 RRF · CrossEncoder 精排+来源多样性", GREEN),
        ("存储与模型层", "Qdrant · PostgreSQL · Neo4j · 本地模型", "向量(617点) / 业务数据(历史中标·审计) / 知识图谱 | BGE-M3 · bge-reranker-v2-m3 · DeepSeek/GLM", GRAY),
    ]
    y = Inches(1.15)
    rh, gap = Inches(1.12), Inches(0.06)
    for name, sub, desc, color in rows:
        add_box(slide, Inches(0.45), y, Inches(2.1), rh, fill=color)
        tb = slide.shapes.add_textbox(Inches(0.5), y + Inches(0.12), Inches(2.0), rh - Inches(0.2))
        _set_text(tb.text_frame, [(name, {"size": 15, "bold": True, "color": WHITE}),
                                  (sub, {"size": 9.5, "color": RGBColor(0xE8, 0xEE, 0xF4)})],
                  align=PP_ALIGN.CENTER)
        add_box(slide, Inches(2.65), y, Inches(10.25), rh, fill=LIGHT)
        tb2 = slide.shapes.add_textbox(Inches(2.85), y + Inches(0.1), Inches(9.9), rh - Inches(0.16))
        _set_text(tb2.text_frame, [(desc, {"size": 12.5, "color": RGBColor(0x33, 0x33, 0x33)})])
        tb2.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        y = Emu(int(y) + int(rh) + int(gap))
    add_text(slide, Inches(0.45), Inches(7.05), Inches(12.4), Inches(0.35),
             [("数据流: 浏览器 → SSE 流式问答 → Agent 编排 → RAG 检索(证据) → 生成(带 [资料N] 引用与来源卡片)",
               {"size": 11.5, "color": TEAL, "bold": True})])


def s_stack(prs):
    slide = slide_base(prs, "二、技术栈")
    groups = [
        ("后端", TEAL, [
            "Python 3.12 + uv (依赖锁定 uv.lock)",
            "FastAPI + Uvicorn (REST/SSE, :8001)",
            "pydantic 配置 + dotenv 环境变量",
            "python-docx/PyMuPDF/openpyxl 文档解析",
            "PaddleOCR/证书 OCR · Fernet 字段加密",
        ]),
        ("前端", PRIMARY, [
            "Next.js 14 (App Router) + React 18",
            "TypeScript + Tailwind CSS",
            "SSE 流式渲染 (fetch stream)",
            "react-markdown + lucide-react",
            "6 页面: 对话/Agent/看板/文档/图谱/资料",
        ]),
        ("存储", GREEN, [
            "Qdrant — 向量库 (hybrid 检索, 617 点)",
            "PostgreSQL 16 — 业务库/历史中标/审计",
            "Neo4j 5 — 知识图谱 (可选)",
            "本地磁盘 / MinIO-S3 证书存储",
            "LRU 内存缓存 (256)",
        ]),
        ("模型与算法", GOLD, [
            "BGE-M3 — Embedding (1024维, Dense+Sparse)",
            "bge-reranker-v2-m3 — CrossEncoder 精排",
            "DeepSeek-V4 / GLM-4.7 — 生成与规划",
            "HF_HOME 本地模型缓存 (D:\\ai_models)",
            "无 GPU 依赖, CPU 可推理",
        ]),
        ("部署与质量", GRAY, [
            "start.bat 原生一键启动 / Docker Compose 5 容器",
            "pytest 分层测试 · 证据门 54 断言",
            "评测: 检索 70 题 · Agent 55 题 · 验收 60 例",
            "audit_logs 双类审计(操作+问答)",
            "JWT/RBAC 三角色 (admin/bidder/匿名)",
        ]),
    ]
    x, y = Inches(0.4), Inches(1.2)
    w, h = Inches(2.49), Inches(5.9)
    for name, color, items in groups:
        add_box(slide, x, y, w, Inches(0.5), fill=color)
        tb = slide.shapes.add_textbox(x, y + Inches(0.05), w, Inches(0.4))
        _set_text(tb.text_frame, [(name, {"size": 14, "bold": True, "color": WHITE})],
                  align=PP_ALIGN.CENTER)
        add_box(slide, x, y + Inches(0.5), w, h - Inches(0.5), fill=LIGHT)
        tb2 = slide.shapes.add_textbox(x + Inches(0.12), y + Inches(0.62),
                                       w - Inches(0.24), h - Inches(0.74))
        _set_text(tb2.text_frame, [("• " + it, {"size": 10.5, "space_after": 7}) for it in items])
        x = Emu(int(x) + int(w) + Inches(0.045))


def s_rag_algo(prs):
    slide = slide_base(prs, "三、核心算法 — RAG 检索链路 (六环节)")
    steps = [
        ("1 文档解析", "PDF 按页/DOCX/扫描件 OCR/Excel\n保留真实页码, 不伪造"),
        ("2 语义切片", "页感知分片, 不切断条款\n标题+正文组合加权 embedding"),
        ("3 Query 规划", "指代消解(会话补全)\n同义词受控变体 ≤3"),
        ("4 多路召回", "BGE-M3 Dense + BM25 Sparse\nQdrant hybrid + RRF(k=60)"),
        ("5 精排", "bge-reranker-v2-m3 重判\n来源多样性 max_per_source"),
        ("6 证据上下文", "证据门: 分数阈值+特征词覆盖\n无引用不生成"),
    ]
    x, y = Inches(0.4), Inches(1.35)
    w, h = Inches(1.98), Inches(2.15)
    for i, (name, desc) in enumerate(steps):
        color = [TEAL, PRIMARY, GOLD, GREEN, PRIMARY, TEAL][i]
        add_box(slide, x, y, w, Inches(0.55), fill=color)
        tb = slide.shapes.add_textbox(x, y + Inches(0.06), w, Inches(0.45))
        _set_text(tb.text_frame, [(name, {"size": 13, "bold": True, "color": WHITE})],
                  align=PP_ALIGN.CENTER)
        add_box(slide, x, y + Inches(0.55), w, h - Inches(0.55), fill=LIGHT)
        tb2 = slide.shapes.add_textbox(x + Inches(0.08), y + Inches(0.65),
                                       w - Inches(0.16), h - Inches(0.75))
        _set_text(tb2.text_frame, [(ln, {"size": 10, "space_after": 3}) for ln in desc.split("\n")])
        if i < 5:
            ar = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Emu(int(x) + int(w)),
                                        y + Inches(0.85), Inches(0.09), Inches(0.35))
            ar.fill.solid()
            ar.fill.fore_color.rgb = GOLD
            ar.line.fill.background()
        x = Emu(int(x) + int(w) + Inches(0.09))

    add_box(slide, Inches(0.4), Inches(3.85), Inches(12.5), Inches(0.42), fill=PRIMARY)
    tb = slide.shapes.add_textbox(Inches(0.55), Inches(3.88), Inches(12.2), Inches(0.36))
    _set_text(tb.text_frame, [("消融实证 (同一 17 题, 逐组件叠加):", {"size": 13, "bold": True, "color": WHITE})])
    ab = [
        ("A0 纯 Dense 单路", "nDCG 0.9756", "基线已可达 100% 召回"),
        ("A1 +BM25 混合", "nDCG 0.6862", "RRF 融合稀释前排"),
        ("A2 +受控变体", "nDCG 0.6520", "多路必需精排兜底"),
        ("A3 +精排+多样性", "nDCG 0.9953", "精排贡献 +34.3pp"),
    ]
    x, y = Inches(0.4), Inches(4.45)
    w = Inches(3.06)
    for i, (name, metric, note) in enumerate(ab):
        fill = GREEN if i == 3 else LIGHT
        tcol = WHITE if i == 3 else RGBColor(0x33, 0x33, 0x33)
        add_box(slide, x, y, w, Inches(1.55), fill=fill)
        tb = slide.shapes.add_textbox(x + Inches(0.1), y + Inches(0.12), w - Inches(0.2), Inches(1.35))
        _set_text(tb.text_frame, [
            (name, {"size": 12, "bold": True, "color": tcol}),
            (metric, {"size": 17, "bold": True, "color": GOLD if i == 3 else PRIMARY}),
            (note, {"size": 10, "color": RGBColor(0xE8, 0xEE, 0xF4) if i == 3 else GRAY}),
        ])
        x = Emu(int(x) + int(w) + Inches(0.085))
    add_text(slide, Inches(0.4), Inches(6.25), Inches(12.5), Inches(0.9), [
        ("扩量复评 (70 题, top-K=5): HitRate@5 = 100% · MRR = 1.0 · nDCG = 0.9933 · EvidenceCoverage = 100%",
         {"size": 13, "bold": True, "color": PRIMARY}),
        ("结论: 覆盖、排序、延迟可控平衡 — 不是简单扩大候选池, 而是每个组件的取舍都有数据支撑",
         {"size": 11.5, "color": GRAY}),
    ])


def s_agent(prs):
    slide = slide_base(prs, "四、核心算法 — Agent 编排与防幻觉")
    add_box(slide, Inches(0.4), Inches(1.15), Inches(6.1), Inches(3.3), fill=LIGHT)
    add_text(slide, Inches(0.6), Inches(1.3), Inches(5.7), Inches(0.4),
             [("ReAct 循环 (主链路, SSE 流式)", {"size": 15, "bold": True, "color": PRIMARY})])
    add_text(slide, Inches(0.6), Inches(1.75), Inches(5.7), Inches(2.6), [
        ("思考 → 行动 → 观察 循环, 最多 3 步", {}),
        ("9 个工具: 知识库/图谱/PG/标书/草稿/异常/范本/投诉/引导", {}),
        ("工具级并行 run_parallel 加速", {}),
        ("复杂多跳: gated 前自动拆分重试一次", {}),
        ("执行过程 SSE 推送, 前端时间线可视化", {}),
    ])

    add_box(slide, Inches(6.8), Inches(1.15), Inches(6.1), Inches(3.3), fill=LIGHT)
    add_text(slide, Inches(7.0), Inches(1.3), Inches(5.7), Inches(0.4),
             [("多专家协作 (multi_agent, 可选入口)", {"size": 15, "bold": True, "color": GREEN})])
    add_text(slide, Inches(7.0), Inches(1.75), Inches(5.7), Inches(2.6), [
        ("主管 LLM 拆解问题 → 调度三专家", {}),
        ("法规 / 案例 / 价格 专家并行 ReAct", {}),
        ("综合终稿: 更全面, 非流式 (1-4 分钟)", {}),
        ("普通问答仍走单 Agent 流式链路", {}),
        ("终稿同样强制事实 + [资料N] 引用", {}),
    ])

    add_box(slide, Inches(0.4), Inches(4.65), Inches(6.1), Inches(2.5), fill=LIGHT)
    add_text(slide, Inches(0.6), Inches(4.78), Inches(5.7), Inches(0.4),
             [("意图路由 — 三业务线显式分类", {"size": 15, "bold": True, "color": GOLD})])
    add_text(slide, Inches(0.6), Inches(5.22), Inches(5.7), Inches(1.85), [
        ("强/弱信号词评分 → 招投标 / 企业 / 法规 / 通用", {}),
        ("法规 9→3 工具 · 企业 9→7 · 保守回退全集", {}),
        ("跨域底座 search_bidding_knowledge 恒保留", {}),
        ("INTENT_ROUTING_ENABLED 开关可回退", {}),
    ])

    add_box(slide, Inches(6.8), Inches(4.65), Inches(6.1), Inches(2.5), fill=LIGHT)
    add_text(slide, Inches(7.0), Inches(4.78), Inches(5.7), Inches(0.4),
             [("证据门 — 硬闸门防幻觉 (54 断言)", {"size": 15, "bold": True, "color": TEAL})])
    add_text(slide, Inches(7.0), Inches(5.22), Inches(5.7), Inches(1.85), [
        ("低分噪声不算证据, 通用 FAQ 不为虚构背书", {}),
        ("特征词 2-gram 覆盖校验, 拦截目录泛化文本", {}),
        ("无证据 → 固定话术拒答, sources=[]", {}),
        ("GATE-01 虚构题 60/60 验收持续 gated", {}),
    ])


def s_fe_chat(prs):
    slide = slide_base(prs, "五、前端展示 — 对话问答页 (SSE 流式 + 来源卡片)")
    _ = add_image_fit(slide, SHOT / "06_chat_sources.png", Inches(0.5),
                      Inches(1.15), Inches(8.6), Inches(6.0))

    items = [
        ("左侧栏: 历史会话 + 功能导航", 0.15, 0.19),
        ("回答区: SSE 流式逐字渲染", 0.60, 0.20),
        ("[资料N] 引用标注可追溯", 0.45, 0.335),
        ("来源卡片: 文件名 + 相关度", 0.50, 0.52),
        ("回答耗时透明 (3.5s)", 0.96, 0.755),
        ("输入框: 联网搜索 / 深度思考", 0.50, 0.88),
        ("模型可切换 (DeepSeek / GLM)", 0.95, 0.04),
    ]
    y = Inches(1.28)
    for text, fx, fy in items:
        add_text(slide, Inches(9.35), y, Inches(3.6), Inches(0.44),
                 [("▸ " + text, {"size": 12, "bold": True, "color": PRIMARY})])
        y = Emu(int(y) + Inches(0.60))
    add_text(slide, Inches(9.35), y + Inches(0.12), Inches(3.6), Inches(1.4), [
        ("无证据 → 固定话术拒答", {"size": 11, "color": GRAY}),
        ("不编造、不给无出处结论", {"size": 11, "color": GRAY}),
        ("问答审计仅存统计量, 不存原文", {"size": 11, "color": GRAY}),
    ])


def s_fe_agent(prs):
    slide = slide_base(prs, "六、前端展示 — 多专家协作过程可视化")
    _ = add_image_fit(slide, SHOT / "v23_multi_agent.png", Inches(0.5),
                      Inches(1.15), Inches(8.6), Inches(6.0))

    items = [
        ("用户问题: 多实体对比", 0.70, 0.145),
        ("输入框 \"多专家协作\" 开关入口", 0.59, 0.958),
        ("主管拆解 → 法规/价格 专家标签", 0.48, 0.29),
        ("专家并行: 轮次/耗时/结论", 0.55, 0.50),
        ("案例专家: 已核查事实 + 来源", 0.55, 0.735),
        ("执行过程透明可追溯", 0.35, 0.233),
    ]
    y = Inches(1.28)
    for text, fx, fy in items:
        add_text(slide, Inches(9.35), y, Inches(3.6), Inches(0.44),
                 [("▸ " + text, {"size": 12, "bold": True, "color": PRIMARY})])
        y = Emu(int(y) + Inches(0.66))
    add_text(slide, Inches(9.35), y + Inches(0.15), Inches(3.6), Inches(1.6), [
        ("评测看板 (Dashboard)", {"size": 13, "bold": True, "color": GREEN}),
        ("检索/Agent 指标可视化", {"size": 11, "color": GRAY}),
        ("HitRate / MRR / nDCG / 通过率", {"size": 11, "color": GRAY}),
    ])


_TOOL_PAGES = [
    ("七、前端展示 — 招投标专项工具 (1/2): 评估分析", [
        ("07_price_panel.png", "价格分析面板",
         ["投标报价对比分析, 异常报价提示", "总价/单价构成透明可查"]),
        ("08_bid_parse_panel.png", "标书解析 + 逐条响应矩阵",
         ["招标要求 vs 投标响应逐条对照", "响应/偏离自动判定, 偏差项高亮"]),
    ]),
    ("七、前端展示 — 招投标专项工具 (2/2): 核验风控", [
        ("v15_cert_ocr.png", "企业资料 + 证书 OCR",
         ["营业执照/资质证书自动识别录入", "敏感字段 Fernet 加密入库"]),
        ("09_collusion_dialog.png", "围串标检测",
         ["两份投标文件相似度对比", "特征码/报价规律等疑点自动提示"]),
    ]),
]


def s_fe_tools(prs):
    for title, imgs in _TOOL_PAGES:
        slide = slide_base(prs, title)
        for k, (fname, cap, descs) in enumerate(imgs):
            x = Inches(0.4) if k == 0 else Inches(6.75)
            _ = add_image_fit(slide, SHOT / fname, x, Inches(1.2),
                              Inches(6.2), Inches(4.45))
            add_box(slide, x, Inches(5.8), Inches(6.2), Inches(1.4), fill=LIGHT)
            tb = slide.shapes.add_textbox(x + Inches(0.18), Inches(5.92),
                                          Inches(5.85), Inches(1.2))
            _set_text(tb.text_frame,
                      [(cap, {"size": 15, "bold": True, "color": PRIMARY})] +
                      [("• " + d, {"size": 12, "color": RGBColor(0x33, 0x33, 0x33),
                                    "space_after": 5}) for d in descs])


def s_quality(prs):
    slide = slide_base(prs, "八、质量保障 — 评测体系与验收数据")
    rows = [
        ("检索评测", "70 题 × 6 指标", "HitRate@5 = 100% · MRR = 1.0 · nDCG = 0.9933 · EvidenceCoverage = 100%", GREEN),
        ("消融对比", "17 题 × 4 档", "A0 纯 Dense 0.9756 → A3 完整 0.9953; 精排单贡献 +34.3pp", GREEN),
        ("Agent 端到端", "55 题真实 HTTP", "通过率 85.5% · 工具选择 87.3% · 事实覆盖 92.6%", TEAL),
        ("HTTP 验收", "60 例全场景", "60/60 通过 · 链路可用性 100% · GATE-01 持续拒答", PRIMARY),
        ("硬闸门", "54 项断言", "无引用不生成 · 低分噪声拦截 · 特征词覆盖校验", GOLD),
        ("审计与安全", "audit_logs", "敏感操作 + 问答交互双审计 · JWT/RBAC 三角色 · Fernet 字段加密", GRAY),
    ]
    y = Inches(1.2)
    for name, scale, result, color in rows:
        add_box(slide, Inches(0.45), y, Inches(1.9), Inches(0.82), fill=color)
        tb = slide.shapes.add_textbox(Inches(0.5), y + Inches(0.08), Inches(1.8), Inches(0.66))
        _set_text(tb.text_frame, [(name, {"size": 13, "bold": True, "color": WHITE})],
                  align=PP_ALIGN.CENTER)
        add_box(slide, Inches(2.4), y, Inches(1.95), Inches(0.82), fill=LIGHT)
        tb2 = slide.shapes.add_textbox(Inches(2.5), y + Inches(0.16), Inches(1.8), Inches(0.55))
        _set_text(tb2.text_frame, [(scale, {"size": 11.5, "bold": True, "color": PRIMARY})],
                  align=PP_ALIGN.CENTER)
        add_box(slide, Inches(4.4), y, Inches(8.5), Inches(0.82), fill=LIGHT)
        tb3 = slide.shapes.add_textbox(Inches(4.6), y + Inches(0.16), Inches(8.15), Inches(0.55))
        _set_text(tb3.text_frame, [(result, {"size": 12.5, "color": RGBColor(0x33, 0x33, 0x33)})])
        y = Emu(int(y) + Inches(0.94))
    add_text(slide, Inches(0.45), Inches(6.95), Inches(12.4), Inches(0.4),
             [("原则: 检索与生成分离评测 · 固定语料固定题集 · 指标可复现可追溯 · 全部离线毫秒级可回归",
               {"size": 12, "bold": True, "color": TEAL})])


def s_deploy(prs):
    slide = slide_base(prs, "九、本地部署 — 两种方式一键启动")
    add_box(slide, Inches(0.4), Inches(1.2), Inches(6.1), Inches(4.9), fill=LIGHT)
    add_text(slide, Inches(0.65), Inches(1.4), Inches(5.6), Inches(0.45),
             [("方式 A: 原生部署 (开发常用)", {"size": 16, "bold": True, "color": PRIMARY})])
    add_text(slide, Inches(0.65), Inches(1.95), Inches(5.6), Inches(4.0), [
        ("1. uv sync — 建 .venv 装依赖", {}),
        ("2. 配置 .env — API Key / Qdrant / 加密密钥", {}),
        ("3. python main.py ingest — Excel 灌知识库", {}),
        ("4. 种子脚本 — PG 历史中标 30 条", {}),
        ("5. start.bat — 自动拉起前后端两个窗口", {}),
        ("", {}),
        ("前端 :3000 · 后端 :8001 · /docs 在线文档", {"bold": True}),
        ("stop.bat 四级兜底停止 (标题/端口/进程树/进程名)", {}),
    ])
    add_box(slide, Inches(6.8), Inches(1.2), Inches(6.1), Inches(4.9), fill=LIGHT)
    add_text(slide, Inches(7.05), Inches(1.4), Inches(5.6), Inches(0.45),
             [("方式 B: Docker Compose 全栈 (交付演示)", {"size": 16, "bold": True, "color": GREEN})])
    add_text(slide, Inches(7.05), Inches(1.95), Inches(5.6), Inches(4.0), [
        ("start-docker.bat 一条命令起 5 容器:", {}),
        ("  backend (python:3.12-slim + uv)", {}),
        ("  frontend (Next.js standalone 产物)", {}),
        ("  qdrant :6333 (volume 持久化)", {}),
        ("  postgres :5432 (volume 持久化)", {}),
        ("  neo4j :7474/7687 (volume 持久化)", {}),
        ("", {}),
        ("健康检查轮询 60s → 自动开浏览器", {"bold": True}),
        ("容器内服务名互联, restart: unless-stopped", {}),
    ])
    add_text(slide, Inches(0.4), Inches(6.4), Inches(12.5), Inches(0.7), [
        ("组件清单: FastAPI API + Next.js 前端 + Qdrant 向量库 + PostgreSQL 业务库 + Neo4j 图谱 + BGE-M3/reranker 本地模型",
         {"size": 12, "bold": True, "color": PRIMARY}),
        ("模型本地化: HF_HOME 重定向 D 盘, 离线可加载, 无 GPU 依赖", {"size": 11, "color": GRAY}),
    ])


def s_end(prs):
    slide = slide_base(prs)
    add_box(slide, 0, 0, SLIDE_W, SLIDE_H, fill=PRIMARY, radius=False)
    add_text(slide, Inches(1), Inches(2.6), Inches(11.3), Inches(1.0),
             [("从 \"能回答\" 到 \"有依据地回答\"", {"size": 36, "bold": True, "color": WHITE})])
    add_text(slide, Inches(1), Inches(3.8), Inches(11.3), Inches(1.2), [
        ("业务价值并非回答更多, 而是回答时更有可靠依据", {"size": 18, "color": RGBColor(0xBD, 0xD7, 0xEE)}),
        ("准确性基础 · 幻觉风险降低 · 可信问答增强", {"size": 15, "color": RGBColor(0xBD, 0xD7, 0xEE)}),
    ])


def main():
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    s_cover(prs)
    s_arch(prs)
    s_stack(prs)
    s_rag_algo(prs)
    s_agent(prs)
    s_fe_chat(prs)
    s_fe_agent(prs)
    s_fe_tools(prs)
    s_quality(prs)
    s_deploy(prs)
    s_end(prs)
    prs.save(OUT)
    print(f"已生成: {OUT} ({len(prs.slides._sldIdLst)} 页)")


if __name__ == "__main__":
    sys.exit(main())
