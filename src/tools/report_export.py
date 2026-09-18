"""报告导出 — Markdown → Word (.docx).

支持:
  - 标题 (## / ### / ####)
  - 粗体 **text**
  - 有序列表 / 无序列表
  - 表格 (Markdown table → Word table)
  - 普通段落

不依赖其他包, 纯 python-docx.
"""
from __future__ import annotations

import re
import tempfile
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import RGBColor

RED = RGBColor(0xFF, 0x00, 0x00)
AMBER = RGBColor(0xC0, 0x60, 0x00)


def _color_of(text: str):
    if "🔴" in text:
        return RED
    if "🟡" in text:
        return AMBER
    return None


def _add_paragraph(doc: Document, text: str, style: str | None = None, bold: bool = False) -> None:
    p = doc.add_paragraph(style=style)
    color = _color_of(text)
    if bold:
        run = p.add_run(text)
        run.bold = True
        if color:
            run.font.color.rgb = color
    else:
        # 处理 **bold** 内联
        parts = re.split(r"(\*\*[^*]+\*\*)", text)
        for part in parts:
            if part.startswith("**") and part.endswith("**"):
                run = p.add_run(part[2:-2])
                run.bold = True
            elif part:
                run = p.add_run(part)
            else:
                continue
            if color:
                run.font.color.rgb = color
                if "🔴" in text:
                    run.bold = True


def _add_markdown_table(doc: Document, header_row: list[str], data_rows: list[list[str]]) -> None:
    table = doc.add_table(rows=1 + len(data_rows), cols=len(header_row))
    table.style = "Table Grid"
    # header
    hdr_cells = table.rows[0].cells
    for i, h in enumerate(header_row):
        hdr_cells[i].text = ""
        run = hdr_cells[i].paragraphs[0].add_run(h.strip())
        run.bold = True
    # data: 含 🔴 的行整行红色加粗, 含 🟡 的行琥珀色
    for ri, row in enumerate(data_rows):
        row_color = RED if any("🔴" in c for c in row) else (
            AMBER if any("🟡" in c for c in row) else None)
        cells = table.rows[ri + 1].cells
        for ci, val in enumerate(row):
            if ci < len(cells):
                cells[ci].text = ""
                run = cells[ci].paragraphs[0].add_run(val.strip())
                if row_color:
                    run.font.color.rgb = row_color
                    if row_color == RED:
                        run.bold = True


def markdown_to_docx(md_text: str, title: str = "报告") -> bytes:
    """Markdown 文本 → Word 文件二进制.

    Args:
        md_text: Markdown 格式的内容
        title: 文档标题 (写进 Word 首页)

    Returns:
        .docx 文件的 bytes (可直接作为 HTTP 响应下载)
    """
    doc = Document()

    # 标题
    heading = doc.add_heading(title, level=0)
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER

    lines = md_text.split("\n")
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()

        # 跳过空行和分隔线
        if not line or line in ("---", "===", "***"):
            i += 1
            continue

        # Markdown 表格检测
        if line.startswith("|") and i + 1 < len(lines) and re.match(r"^\|[\s:\-|]+\|$", lines[i + 1]):
            header = [c.strip() for c in line.strip("|").split("|")]
            i += 2  # 跳过分隔线
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                rows.append([c.strip() for c in lines[i].strip("|").split("|")])
                i += 1
            _add_markdown_table(doc, header, rows)
            continue

        # 标题
        h_match = re.match(r"^(#{1,4})\s+(.+)", line)
        if h_match:
            level = min(len(h_match.group(1)), 4)
            doc.add_heading(h_match.group(2).strip(), level=level)
            i += 1
            continue

        # 无序列表
        list_match = re.match(r"^[-*+]\s+(.+)", line)
        if list_match:
            _add_paragraph(doc, list_match.group(1), style="List Bullet")
            i += 1
            continue

        # 有序列表
        ol_match = re.match(r"^\d+\.\s+(.+)", line)
        if ol_match:
            _add_paragraph(doc, ol_match.group(1), style="List Number")
            i += 1
            continue

        # 普通段落
        _add_paragraph(doc, line)
        i += 1

    # 写入临时文件再读 bytes
    with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as tmp:
        tmp_path = tmp.name
    doc.save(tmp_path)
    try:
        data = Path(tmp_path).read_bytes()
    finally:
        Path(tmp_path).unlink(missing_ok=True)
    return data


def export_bid_to_docx(markdown: str, project_name: str = "投标书草稿") -> bytes:
    """投标书专用导出: 标题用项目名."""
    title = project_name or "投标书草稿"
    return markdown_to_docx(markdown, title=title)
